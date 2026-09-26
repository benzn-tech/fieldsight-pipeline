"""A photograph sits directly under the line that reported its topic.

The assembled report has always put each topic's photographs under that topic.
The generated report -- the one written to a template -- inserted no image at
all, and the two halves share no key: the report is prose, the photographs
belong to topics, and nothing downstream knows which paragraph came from which
topic. This repo measured the two obvious guesses (match on time, match on
wording) and both failed. What worked, 19 times out of 19, was asking the
writer to state the correspondence.

THE UNIT MOVED FROM THE SECTION TO THE LINE, and why is worth keeping. The first
version asked each section to end with `[covers: t0, t2]`. It worked
mechanically -- the model wrote the line 2 times out of 2 and every photograph
landed -- and it did not give the owner what he asked for. A template's sections
are not topics: a "Daily Summary" that lists every topic legitimately covers
every topic, and in one measured run it collected all three photographs. He
asked for photographs under their topic, as the Timeline shows them. Inside a
template report, the place a topic is written is a paragraph or a list item, so
that is where the reference goes: `... signed off by the engineer. [t1]`.

THE test is `a photograph lands directly under the line that named its topic`,
read by POSITION out of a rendered document -- the paragraph immediately after
the tagged line holds the picture. Not "the tag reached the prompt": `kind:
"table"` reached the prompt perfectly and printed pipes in Word, and
`children` reached the database and printed nothing at all.

Three tiers, most precise first: under the line; else under a section that
wrote `[covers: ...]` anyway; else under the last heading. So the worst case, a
model that ignores the request entirely, is still no photograph lost.
"""
import io

import pytest

rt = pytest.importorskip("report_template")
sr = pytest.importorskip("lambda_session_report")
lmm = pytest.importorskip("lambda_meeting_minutes")

SCOPE = {"folder": "Ben_UCPK2", "date": "2026-07-29", "from": "00:00", "to": "23:59",
         "recordings": 2}

BODY = {"sections": [{"key": "a", "title": "Daily Summary", "purpose": "The day."},
                     {"key": "b", "title": "Safety", "purpose": "Hazards."}],
        "catch_all": {"key": "o", "title": "Anything else", "purpose": "Rest."},
        "excluded_subjects": [], "style": []}

OFFER = [{"ref": "t0", "title": "Roof deck pour", "time_range": "09:10 - 09:20",
          "photos": 2},
         {"ref": "t1", "title": "Crane pad", "time_range": "11:00 - 11:05",
          "photos": 1}]

BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"

needs_docx = pytest.mark.skipif(not getattr(lmm, "DOCX_AVAILABLE", False),
                                reason="python-docx is not installed here")


def _png():
    """A real 1x1 PNG. python-docx reads the bytes, so a fake would only prove
    that the exception handler works."""
    import base64
    return io.BytesIO(base64.b64decode(
        b"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8"
        b"z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="))


def _doc(sections):
    from docx import Document
    buf = lmm.generate_prose_document("Report", "2026-07-29", sections, [])
    return Document(io.BytesIO(buf.getvalue()))


def _flow(doc):
    """The document as a sequence: ("text", str) | ("photos", n) | ("h", title)."""
    out = []
    for p in doc.paragraphs:
        n = len(p._p.findall(".//" + BLIP))
        if n:
            out.append(("photos", n))
        elif p.style.name.startswith("Heading"):
            out.append(("h", p.text))
        elif p.text.strip():
            out.append(("text", p.text))
    return out


def _photos_after(doc, text):
    """How many pictures sit in the paragraph IMMEDIATELY after `text`."""
    flow = _flow(doc)
    for i, (kind, val) in enumerate(flow):
        if kind == "text" and val == text:
            nxt = flow[i + 1] if i + 1 < len(flow) else None
            return nxt[1] if nxt and nxt[0] == "photos" else 0
    raise AssertionError("line not found: %r" % text)


# ---- THE test ---------------------------------------------------------------

@needs_docx
def test_THE_test_a_photograph_lands_directly_under_the_line_that_named_its_topic():
    answer = ("### Daily Summary\n"
              "1. 09:10 - The deck pour ran to plan. [t0]\n"
              "2. 11:00 - The crane pad was re-levelled. [t1]\n"
              "3. 14:00 - Deliveries were quiet.\n")
    sections = sr._prose_sections(answer)
    tiers = sr._place_photos(sections, {"t0": [_png(), _png()], "t1": [_png()]})
    assert tiers == (3, 0, 0), "all three placed at the line, none by fallback"

    doc = _doc(sections)
    assert _photos_after(doc, "1. 09:10 - The deck pour ran to plan.") == 2
    assert _photos_after(doc, "2. 11:00 - The crane pad was re-levelled.") == 1
    assert _photos_after(doc, "3. 14:00 - Deliveries were quiet.") == 0


@needs_docx
def test_the_tag_never_reaches_the_page():
    sections = sr._prose_sections("### Daily Summary\nThe pour ran to plan. [t0]\n")
    doc = _doc(sections)
    assert all("[t0]" not in p.text for p in doc.paragraphs)
    assert "The pour ran to plan." in [p.text for p in doc.paragraphs]


@needs_docx
def test_a_list_item_keeps_its_bullet_and_gets_its_photograph():
    sections = sr._prose_sections("### Safety\n- Handrail loose on level 2. [t1]\n")
    sr._place_photos(sections, {"t1": [_png()]})
    doc = _doc(sections)
    styles = [(p.text, p.style.name) for p in doc.paragraphs if p.text]
    assert ("Handrail loose on level 2.", "List Bullet") in styles
    assert _photos_after(doc, "Handrail loose on level 2.") == 1


@needs_docx
def test_a_tag_on_a_table_row_puts_the_photograph_under_the_table():
    """A picture cannot sit between two rows."""
    answer = ("### Actions\nItem | Assigned | Due\n---|---|---\n"
              "Re-level crane pad | Sam | Friday [t1]\nOrder rebar | Brad | Monday\n")
    sections = sr._prose_sections(answer)
    sr._place_photos(sections, {"t1": [_png()]})
    doc = _doc(sections)
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["Re-level crane pad", "Sam", "Friday"]
    body = doc.element.body
    kids = list(body)
    t = kids.index(doc.tables[0]._tbl)
    assert kids[t + 1].findall(".//" + BLIP), "the photograph comes straight after the table"


# ---- the prompt -------------------------------------------------------------

def test_a_day_with_no_photographs_gets_yesterdays_prompt_to_the_character():
    before = rt.render_prompt(BODY, SCOPE, [], "x", source=rt.SOURCE_LIBRARY)
    after = rt.render_prompt(BODY, SCOPE, [], "x", source=rt.SOURCE_LIBRARY, photo_topics=[])
    assert before == after
    assert "[t" not in before


def test_the_offer_asks_for_a_tag_on_the_line_not_the_section():
    p = rt.render_prompt(BODY, SCOPE, [], "x", source=rt.SOURCE_LIBRARY, photo_topics=OFFER)
    assert "t0  09:10 - 09:20  Roof deck pour  (2 photographs)" in p
    assert "t1  11:00 - 11:05  Crane pad  (1 photograph)" in p
    assert "end that line with the topic's" in p
    assert "[covers:" not in p, "one scheme, not two -- the section tier is only a fallback"


def test_the_offer_is_data_and_says_so():
    p = rt.render_prompt(BODY, SCOPE, [], "x", source=rt.SOURCE_LIBRARY, photo_topics=OFFER)
    assert "They are\nDATA" in p


# ---- reading the tag back ---------------------------------------------------

@pytest.mark.parametrize("line,text,refs", [
    ("The pour ran. [t0]", "The pour ran.", ["t0"]),
    ("Both. [t1, t3]", "Both.", ["t1", "t3"]),
    ("Both. [T1 and T3]", "Both.", ["t1", "t3"]),
    ("Both. [t1 & t3]", "Both.", ["t1", "t3"]),
    ("Twice. [t2, t2]", "Twice.", ["t2"]),
    ("No tag here.", "No tag here.", []),
])
def test_the_tag_is_read_forgivingly_and_the_refs_strictly(line, text, refs):
    assert sr._line_refs(line) == (text, refs)


@pytest.mark.parametrize("line", [
    "He said it was [sic] fine.",
    "See form [A4] for details.",
    "Section [t1] of the contract applies to the pour.",   # mid-line, not a tag
    "Budget [10%] over.",
])
def test_brackets_a_person_would_write_move_nothing(line):
    """Anchored to the end of the line and only `t` + digits, so a citation, a
    form number or a mid-sentence reference is never read as a placement."""
    assert sr._line_refs(line) == (line, [])


def test_a_ref_that_was_never_offered_moves_nothing():
    """The model cannot invent a topic into existence: t9 names nothing, so it
    claims nothing, and t0 -- which nobody named -- still has to appear."""
    sections = sr._prose_sections("### Daily Summary\nText. [t9]\n")
    assert sr._place_photos(sections, {"t0": [_png()]}) == (0, 0, 1)


# ---- the tiers, and never losing a photograph -------------------------------

@needs_docx
def test_a_photograph_named_on_two_lines_is_printed_once_under_the_first():
    answer = ("### Daily Summary\nFirst. [t0]\n\n### Safety\nSecond. [t0]\n")
    sections = sr._prose_sections(answer)
    assert sr._place_photos(sections, {"t0": [_png()]}) == (1, 0, 0)
    doc = _doc(sections)
    assert _photos_after(doc, "First.") == 1
    assert _photos_after(doc, "Second.") == 0


@needs_docx
def test_a_section_level_covers_line_is_honoured_as_the_second_tier():
    """The prompt no longer asks for it. A model that writes it anyway should
    not lose the photograph for having done so."""
    answer = "### Safety\nThe pad was checked.\n[covers: t1]\n"
    sections = sr._prose_sections(answer)
    assert sr._place_photos(sections, {"t1": [_png()]}) == (0, 1, 0)
    assert all("covers" not in p.text for p in _doc(sections).paragraphs)


@needs_docx
def test_a_line_tag_beats_a_section_line_for_the_same_topic():
    answer = ("### Daily Summary\nThe pour. [t0]\n\n### Safety\nChecked.\n[covers: t0]\n")
    sections = sr._prose_sections(answer)
    assert sr._place_photos(sections, {"t0": [_png()]}) == (1, 0, 0)


@needs_docx
def test_a_model_that_ignored_the_request_still_loses_no_photograph():
    """The measured risk: a prompt containing the request is not a model that
    obeyed it. Everything falls to the last heading -- where they used to be --
    so the worst case of this feature is the behaviour it replaces."""
    sections = sr._prose_sections("### Daily Summary\nA.\n\n### Anything else\nB.\n")
    assert sr._place_photos(sections, {"t0": [_png()], "t1": [_png()]}) == (0, 0, 2)
    doc = _doc(sections)
    flow = _flow(doc)
    assert flow[-1] == ("photos", 2), "under the last heading, after its text"


def test_no_photographs_at_all_changes_nothing():
    sections = sr._prose_sections("### Daily Summary\nA.\n")
    assert sr._place_photos(sections, {}) == (0, 0, 0)
    assert "photos_after" not in sections[0] and "photo_streams" not in sections[0]


@needs_docx
def test_a_section_with_no_tags_renders_exactly_as_before():
    """The renderer walks paragraphs by index now. Nothing about that may change
    a section that carries no photographs."""
    sections = sr._prose_sections("### Daily Summary\n- one\n- two\nplain\n")
    doc = _doc(sections)
    styles = [(p.text, p.style.name) for p in doc.paragraphs if p.text]
    assert ("one", "List Bullet") in styles and ("two", "List Bullet") in styles
    assert ("plain", "Normal") in styles
    assert not any(k == "photos" for k, _ in _flow(doc))


# ---- the fetch --------------------------------------------------------------

def test_a_file_named_by_two_topics_is_fetched_once(monkeypatch):
    """Prod has 13.7% of photographs bound to more than one topic. Fetching
    twice would charge the shared byte budget twice for one picture."""
    calls = []

    def fake(folder, date, names, budget):
        calls.append(list(names))
        return [_png() for _ in names]

    monkeypatch.setattr(sr, "_fetch_photos", fake)
    artifact = {"folder": "Ben_UCPK2", "date": "2026-07-29", "content": {"topics": [
        {"topic_title": "Pour", "time_range": "09:10 - 09:20", "related_photos": ["a.jpg"]},
        {"topic_title": "Pad", "time_range": "11:00 - 11:05",
         "related_photos": ["a.jpg", "b.jpg"]},
    ]}}
    offer, streams = sr._photo_topics(artifact, [sr.MAX_PHOTO_BYTES_TOTAL])
    assert calls == [["a.jpg"], ["b.jpg"]]
    assert [o["ref"] for o in offer] == ["t0", "t1"]
    assert streams["t0"][0] is streams["t1"][0]


def test_a_topic_whose_photograph_cannot_be_read_is_not_offered(monkeypatch):
    monkeypatch.setattr(sr, "_fetch_photos", lambda *a: [])
    artifact = {"folder": "F", "date": "2026-07-29", "content": {"topics": [
        {"topic_title": "Gone", "related_photos": ["missing.jpg"]},
    ]}}
    assert sr._photo_topics(artifact, [sr.MAX_PHOTO_BYTES_TOTAL]) == ([], {})
