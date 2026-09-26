"""Choosing "Table" produces a table, or the choice is not offered.

`kind` reached the editor, the body, the database and the worker, and
`render_prompt` never read it -- while the editor's own Test render DID honour
it. So a person who chose "List" saw a list in the preview and got prose in the
document, and the only way to get the format they had already asked for was to
write it into the description in words. That free text then went into the prompt
beside our rules, which is the thing the rest of this change is about.

Wiring `kind` up is therefore not a nicety; it is the supply side. But wiring it
only as far as the PROMPT would have moved the lie rather than removed it for
"Table": this renderer knows headings, bullets and paragraphs, so a model asked
for a markdown table produced pipe characters, and pipe characters in a Word
document are worse than the prose they replaced.

THE test is `pipe rows become a table and not a paragraph of pipes`.

`photos` went the other way and is refused at the door: nothing in this path
inserts an image, so there is no version of it that is not a heading with a
sentence underneath. Refusing is the honest half of the same rule.
"""
import pytest

lmm = pytest.importorskip("lambda_meeting_minutes")

if not getattr(lmm, "DOCX_AVAILABLE", False):
    pytest.skip("python-docx is not installed here (it is in CI via the layer)",
                allow_module_level=True)

ROWS = [
    "| Item | Owner | When |",
    "|---|---|---|",
    "| Crane pad | Sam | Friday |",
    "| Site power | Brad | next week |",
]


def _doc(paragraphs):
    return lmm.generate_prose_document(
        "Report", "2026-09-23", [{"title": "Deliveries", "paragraphs": paragraphs}], [])


def _rendered(paragraphs):
    from docx import Document
    import io as _io
    buf = _doc(paragraphs)
    return Document(_io.BytesIO(buf.getvalue()))


# ---- THE test ---------------------------------------------------------------

def test_THE_test_pipe_rows_become_a_table_and_not_a_paragraph_of_pipes():
    doc = _rendered(ROWS)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Item", "Owner", "When"]
    assert [c.text for c in table.rows[1].cells] == ["Crane pad", "Sam", "Friday"]
    assert len(table.rows) == 3, "the |---| rule is a rule, not a row"

    for p in doc.paragraphs:
        assert "|" not in p.text, "a pipe left in the prose is the old behaviour"


def test_text_around_a_table_still_reads_as_text():
    doc = _rendered(["Three deliveries landed."] + ROWS + ["Nothing else moved."])
    texts = [p.text for p in doc.paragraphs]
    assert "Three deliveries landed." in texts
    assert "Nothing else moved." in texts
    assert len(doc.tables) == 1


def test_a_ragged_table_does_not_lose_a_cell_or_raise():
    """The model writes these; it will not always write them square."""
    doc = _rendered(["| A | B | C |", "| only one |"])
    table = doc.tables[0]
    assert len(table.columns) == 3
    assert [c.text for c in table.rows[1].cells] == ["only one", "", ""]


def test_a_sentence_with_a_pipe_in_it_is_not_a_table():
    doc = _rendered(["The board reads Gate A | Gate B."])
    assert not doc.tables
    assert "The board reads Gate A | Gate B." in [p.text for p in doc.paragraphs]


def test_a_one_column_table_comes_out_as_lines():
    """Measured, not imagined: a section that said "table" and named no columns
    got a single column from the model, headed with the section's own title,
    holding lines that had read fine as sentences the day before. The prompt
    names the columns now, so this should not arrive -- and when it does, the
    lines are worth more as lines than as a column of boxes."""
    doc = _rendered(["| Open Actions |",
                     "| **no owner recorded** - Platform login |",
                     "| **no owner recorded** - Onboarding session |"])
    assert not doc.tables, "one column is not a table"
    texts = [p.text for p in doc.paragraphs]
    assert "Open Actions" in texts
    assert "**no owner recorded** - Platform login" in texts
    for p in doc.paragraphs:
        assert "|" not in p.text


def test_two_columns_is_still_a_table():
    """The floor is at one, not at three -- a two-column table is a table."""
    doc = _rendered(["| Item | Due |", "|---|---|", "| Crane pad | Friday |"])
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["Crane pad", "Friday"]


# ---- the outer pipes are optional ------------------------------------------

# Verbatim from a report generated through the ordinary browser flow on TEST,
# 2026-09-26, template "daily report" v10. The renderer printed all four lines
# as paragraphs -- pipes and the dashed rule in the Word document -- because it
# only recognised rows that START with `|`. The outer pipes are optional in
# GitHub-flavoured markdown and the model uses both forms; the reports checked
# before shipping all happened to use outer pipes.
NO_OUTER_PIPES = [
    "Item | Assigned | Due",
    "---|---|---",
    "**no owner recorded** - Platform initial login using temporary password then "
    "change to own password per PDF - *no date* | no owner recorded | no date",
    "**no owner recorded** - Elevation onboarding session to attend tomorrow at "
    "eleven o'clock - *tomorrow morning, eleven o'clock* | no owner recorded | "
    "tomorrow morning, eleven o'clock",
]


def test_a_table_without_outer_pipes_is_still_a_table():
    """THE regression. The delimiter row is the signal, not a leading pipe."""
    doc = _rendered(NO_OUTER_PIPES)
    assert len(doc.tables) == 1, "it was printed as four paragraphs of pipes"
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Item", "Assigned", "Due"]
    assert len(table.rows) == 3, "the ---|---|--- line is a rule, not a row"
    assert [c.text for c in table.rows[1].cells][1:] == ["no owner recorded", "no date"]
    for p in doc.paragraphs:
        assert "|" not in p.text, "a pipe in the prose is the defect itself"
        assert not p.text.startswith("---"), "the delimiter row leaked as text"


def test_aligned_delimiters_count_too():
    doc = _rendered(["Item | Due", ":--- | ---:", "Pour | Friday"])
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["Pour", "Friday"]


def test_text_after_a_pipeless_table_is_not_swallowed_into_it():
    doc = _rendered(["Item | Due", "---|---", "Pour | Friday", "Nothing else moved."])
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 2
    assert "Nothing else moved." in [p.text for p in doc.paragraphs]


def test_two_sentences_with_pipes_are_not_a_table():
    """Without a delimiter row there is no table, however many pipes the prose
    happens to contain -- that is what makes the delimiter a safe signal."""
    doc = _rendered(["The board reads Gate A | Gate B.", "Signage says North | South."])
    assert not doc.tables


def test_a_bare_rule_is_not_a_delimiter():
    """`---` alone is a horizontal rule or a setext underline, not a table."""
    doc = _rendered(["Heading-ish line", "---", "More text."])
    assert not doc.tables


def test_bullets_and_prose_are_unchanged():
    """Nothing above may cost the two things this renderer already did."""
    doc = _rendered(["- first", "* second", "a sentence"])
    styles = [(p.text, p.style.name) for p in doc.paragraphs if p.text]
    assert ("first", "List Bullet") in styles
    assert ("second", "List Bullet") in styles
    assert ("a sentence", "Normal") in styles
