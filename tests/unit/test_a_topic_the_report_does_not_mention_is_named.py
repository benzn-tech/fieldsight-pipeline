"""A recorded topic the report does not mention is named at the end of it.

WHY THIS EXISTS. A section description can tell the model what to leave out,
and the owner decided on 2026-09-27 that descriptions keep that power -- "NO
DATA TODAY" is a feature. Probe 3 measured a description beating the house
rule 10 times out of 10 with the fence in place. So nothing written INTO the
prompt can make an omission visible: the only guard that a description cannot
reach is one built after the answer, from a count, in our code.

WHAT IS COUNTED. Every topic in the window is offered with a ref (`t0`, `t2`),
and the model ends each line that reports one with it -- the mechanism photo
placement already measured at 15/15. A topic no line names is listed under
"Also recorded", after the actions table, with its time.

WHAT IT DOES NOT CLAIM. The refs are the model's own statement. A topic the
model did not write about is caught; a line tagged with a topic it did not
really report is not. The note says "not referred to", which is what was
counted, and never "left out".

THE test is `a topic the answer never names is listed after the actions`,
read by POSITION out of a rendered document -- the note has to come after the
actions table, or a model section titled "Actions" would put its table under
the note's heading.
"""
import datetime as dt
import io

import pytest

import lambda_session_report as sr

lmm = pytest.importorskip("lambda_meeting_minutes")
needs_docx = pytest.mark.skipif(not getattr(lmm, "DOCX_AVAILABLE", False),
                                reason="python-docx is not installed here")

BODY = {
    "sections": [{"key": "sum", "title": "Daily Summary", "purpose": "The day."},
                 {"key": "actions", "title": "Actions", "purpose": "One line each."}],
    "catch_all": {"key": "other", "title": "Anything else", "purpose": "The rest."},
    "excluded_subjects": [], "style": [],
}

TOPICS = [
    {"topic_title": "Platform onboarding", "time_range": "13:20 - 13:22",
     "action_items": [{"action": "Send invites", "responsible": "Ben", "deadline": None}]},
    {"topic_title": "Camera placement", "time_range": "13:38 - 13:39"},
    {"topic_title": "Concrete video recap", "time_range": "13:39 - 13:40"},
]


def _artifact(window=("00:00", "23:59"), topics=TOPICS):
    return {
        "requestId": "r1", "folder": "Ben_UCPK2", "date": "2026-09-23",
        "sessionId": "sid" + "a" * 32, "resultKey": "session_report_results/x.json",
        "title": "Day report", "deliver": "download",
        "generate": {"templateId": "u", "templateVersion": 1, "templateName": "Daily",
                     "templateBody": BODY, "templateSource": "library"},
        "window": {"from": window[0], "to": window[1]},
        "excludedTopics": [],
        "content": {"date": "2026-09-23", "participants": ["Ben"], "topics": list(topics)},
    }


# The answer a description like "do not mention camera placement" produces:
# t1 is simply not there.
ANSWER_WITHOUT_T1 = ("### Daily Summary\nOnboarding for the platform was arranged. [t0]\n"
                     "The concrete video was recapped. [t2]\n\n"
                     "### Actions\n- **Ben** - send invites - *no date*\n\n"
                     "### Anything else\nNothing here.\n")

ANSWER_WITH_ALL = ANSWER_WITHOUT_T1.replace(
    "Nothing here.", "Camera positions were discussed. [t1]")


@pytest.fixture
def run(monkeypatch):
    seen = {}

    def go(answer, artifact=None, render=True):
        monkeypatch.setattr(sr.transcript_window, "select_keys",
                            lambda *a, **k: [(dt.datetime(2026, 9, 23, 13, 20), "t/x.json")])
        monkeypatch.setattr(sr.transcript_window, "assemble",
                            lambda *a, **k: [{"at": dt.datetime(2026, 9, 23, 13, 20),
                                              "line": "[13:20:00] Ben: hello"}])

        def fake_call(prompt, **kw):
            seen["prompt"] = prompt
            return answer, None

        monkeypatch.setattr(sr.llm_utils, "call_llm", fake_call)
        monkeypatch.setattr(sr.llm_utils, "active_model", lambda **k: "m")
        if not render:
            monkeypatch.setattr(sr.lambda_meeting_minutes, "generate_prose_document",
                                lambda *a, **k: io.BytesIO(b"x"))
        buf, meta = sr._generate_document(artifact or _artifact())
        seen["meta"] = meta
        seen["buf"] = buf
        return seen

    return go


def _flow(buf):
    """The document as ("h", title) | ("p", text) | ("table", first_header) in order."""
    from docx import Document
    doc = Document(io.BytesIO(buf.getvalue()))
    paras = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    out = []
    for el in doc.element.body.iterchildren():
        if el in tables:
            out.append(("table", tables[el].rows[0].cells[0].text))
        elif el in paras and paras[el].text.strip():
            p = paras[el]
            out.append(("h" if p.style.name.startswith("Heading") else "p", p.text.strip()))
    return out


# ---- THE test -----------------------------------------------------------------

@needs_docx
def test_THE_test_a_topic_the_answer_never_names_is_listed_after_the_actions(run):
    flow = _flow(run(ANSWER_WITHOUT_T1)["buf"])
    table = flow.index(("table", "Action"))
    note = flow.index(("h", sr.COVERAGE_TITLE))
    assert note > table, "after the actions table, never between it and its heading"
    assert flow[note + 1] == ("p", sr.COVERAGE_INTRO)
    assert flow[note + 2] == ("p", "13:38 - 13:39  Camera placement")
    assert flow[note + 3:] == [], "only the topic nobody named"


@needs_docx
def test_a_report_that_names_every_topic_has_no_note(run):
    flow = _flow(run(ANSWER_WITH_ALL)["buf"])
    assert ("h", sr.COVERAGE_TITLE) not in flow


def test_the_result_records_what_the_note_printed(run):
    meta = run(ANSWER_WITHOUT_T1, render=False)["meta"]
    assert meta["topicsOffered"] == 3
    assert meta["topicsNotReferenced"] == [
        {"ref": "t1", "title": "Camera placement", "time_range": "13:38 - 13:39"}]


def test_every_topic_is_offered_even_without_photographs(run):
    prompt = run(ANSWER_WITH_ALL, render=False)["prompt"]
    for line in ("t0  13:20 - 13:22  Platform onboarding  (no photographs)",
                 "t1  13:38 - 13:39  Camera placement  (no photographs)",
                 "t2  13:39 - 13:40  Concrete video recap  (no photographs)"):
        assert line in prompt


# ---- what counts as naming a topic ----------------------------------------------

def test_a_section_covers_line_counts_as_naming_it(run):
    answer = ANSWER_WITHOUT_T1.replace("Nothing here.", "Cameras.\n[covers: t1]")
    assert run(answer, render=False)["meta"]["topicsNotReferenced"] == []


def test_a_ref_that_was_never_offered_names_nothing(run):
    """t9 is not a topic. Writing it must not make t1 look reported."""
    answer = ANSWER_WITHOUT_T1.replace("Nothing here.", "Something. [t9]")
    missing = run(answer, render=False)["meta"]["topicsNotReferenced"]
    assert [t["ref"] for t in missing] == ["t1"]


def test_a_model_that_tagged_nothing_lists_every_topic(run):
    """The failure mode of the mechanism itself: no tags at all. Every topic is
    listed -- noisy, and honest: nothing was COUNTED as reported."""
    answer = "### Daily Summary\nA day.\n"
    missing = run(answer, render=False)["meta"]["topicsNotReferenced"]
    assert [t["ref"] for t in missing] == ["t0", "t1", "t2"]


# ---- the window --------------------------------------------------------------------

def test_a_topic_outside_the_window_is_not_offered_or_listed(run):
    seen = run("### Daily Summary\nOnboarding. [t0]\n", render=False,
               artifact=_artifact(window=("13:00", "13:30")))
    assert "Camera placement" not in seen["prompt"]
    assert seen["meta"]["topicsOffered"] == 1
    assert seen["meta"]["topicsNotReferenced"] == []


def test_a_topic_with_no_time_is_counted_only_when_the_window_is_the_whole_day(run):
    """It cannot be placed. Across the whole day it is in the window by
    definition; anywhere narrower the note would be guessing."""
    topics = [{"topic_title": "Undated", "time_range": ""}]
    whole = run("### Daily Summary\nA.\n", render=False, artifact=_artifact(topics=topics))
    assert whole["meta"]["topicsOffered"] == 1
    narrow = run("### Daily Summary\nA.\n", render=False,
                 artifact=_artifact(window=("13:00", "13:30"), topics=topics))
    assert narrow["meta"]["topicsOffered"] == 0


@pytest.mark.parametrize("time_range,expected", [
    ("12:50 - 13:05", True),    # straddles the start
    ("13:25 - 13:40", True),    # straddles the end
    ("13:30 - 13:40", False),   # starts where the window ends
    ("12:00 - 13:00", False),   # ends where the window starts
    ("13:10", True),            # a point inside (BUG-09 collapse)
    ("13:00", True),            # a point on the start
])
def test_overlap_decides_the_window(time_range, expected):
    frm, to = sr._clock("2026-09-23", "13:00"), sr._clock("2026-09-23", "13:30")
    assert sr._in_window({"time_range": time_range}, "2026-09-23", frm, to) is expected


# ---- nothing the model writes reaches the note ---------------------------------------

@needs_docx
def test_a_section_the_model_titles_like_the_note_does_not_suppress_it(run):
    """The note is ours. A model heading that happens to read "Also recorded"
    is just a heading; the note still follows the actions."""
    answer = ANSWER_WITHOUT_T1.replace("### Anything else", "### Also recorded")
    flow = _flow(run(answer)["buf"])
    assert [i for i, x in enumerate(flow) if x == ("h", sr.COVERAGE_TITLE)][-1] > \
        flow.index(("table", "Action"))
    assert ("p", "13:38 - 13:39  Camera placement") in flow


@needs_docx
def test_a_photograph_of_a_topic_nobody_named_sits_under_its_note_line(monkeypatch, run):
    import base64
    png = base64.b64decode(
        b"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8"
        b"z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    def fetch(folder, date, names, budget, names_out=None, edge=None):
        if names_out is not None:
            names_out.extend(names)
        return [io.BytesIO(png) for _ in names]
    monkeypatch.setattr(sr, "_fetch_photos", fetch)
    topics = [dict(TOPICS[0]), dict(TOPICS[1], related_photos=["cam.jpg"]), dict(TOPICS[2])]
    seen = run(ANSWER_WITHOUT_T1, artifact=_artifact(topics=topics))
    assert seen["meta"]["photosUnplaced"] == 0
    from docx import Document
    doc = Document(io.BytesIO(seen["buf"].getvalue()))
    blip = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    paras = doc.paragraphs
    i = [k for k, p in enumerate(paras) if p.text.strip() == "13:38 - 13:39  Camera placement"][0]
    assert paras[i + 1]._p.findall(".//" + blip), "the photograph is directly under its line"
