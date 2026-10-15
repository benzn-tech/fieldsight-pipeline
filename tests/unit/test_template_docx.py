"""Reading a checklist out of a Word form -- no model (spec 2026-10-16 section 3).

Synthetic forms mirror the two real layouts measured in the spike:
  BICL   one big table; a column-header row, then single-cell (merged) section rows,
         each followed by item rows whose answer cell reads "Yes No N/A".
  Penina one table per section; the header row's first cell IS the section title
         ("General | Yes/No/NA | Comment | ..."), and filled-in comments sit beside items.
python-docx cannot author w14 checkbox content controls, so the synthetic forms use the
literal answer text. The real forms do NOT pin the content-control path either: every
w14 checkbox cell in them also reads "Yes No N/A" as text (measured 2026-10-16, 47/47 and
115/115), so a bare w14:checkbox element is injected by hand in its own test below.
"""
import json
import os
import pathlib
from io import BytesIO

import pytest

docx = pytest.importorskip("docx")
import template_docx  # noqa: E402

FIXTURES = pathlib.Path(__file__).resolve().parents[1] / "fixtures" / "template_digitise"
REAL = {
    "bicl": os.environ.get("TEMPLATE_DOCX_BICL",
                           "C:/Users/camil/Dropbox/ADE/12.15 - Site Establishment Checklist (BICL).docx"),
    "penina": os.environ.get("TEMPLATE_DOCX_PENINA",
                             "C:/Users/camil/Dropbox/ADE/Penina Site Inspection Checklist.docx"),
}


def _save(d):
    buf = BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf


def _row(table, *cells, merge=False):
    r = table.add_row()
    if merge:
        merged = r.cells[0].merge(r.cells[len(r.cells) - 1])
        merged.text = cells[0]
    else:
        for i, c in enumerate(cells):
            r.cells[i].text = c


def _bicl_like():
    d = docx.Document()
    d.add_paragraph("Site Establishment Checklist")
    details = d.add_table(rows=0, cols=2)
    _row(details, "Project Name", "")
    t = d.add_table(rows=0, cols=4)
    _row(t, "ACTIVITY & REQUIREMENTS", "COMPLIES Y / N / NA", "DUE BY", "COMMENT")
    _row(t, "Site Boundaries", merge=True)
    _row(t, "Site secured", "Yes No N/A", "", "")
    _row(t, "Hoardings established\nwhere appropriate", "\u2610 Yes \u2610 No \u2610 N/A", "", "")
    _row(t, "Emergency Procedures", merge=True)
    _row(t, "Evacuation plan displayed", "Yes No N/A", "", "")
    other = d.add_table(rows=0, cols=1)
    _row(other, "Other Comments")
    return _save(d)


def _penina_like():
    d = docx.Document()
    head = d.add_table(rows=0, cols=2)
    _row(head, "Project: Penina", "Date: 26/1/26")
    for title, items in (("General", ["Site Hazard Identification Board display",
                                      "Is the H&S Policy displayed?"]),
                         ("Site Amenities", ["Is the toilet facility in a clean and sanitary state?"])):
        t = d.add_table(rows=0, cols=5)
        _row(t, title, "Yes/No/NA", "Comment", "Person Responsible", "Target Close Out Date")
        for it in items:
            _row(t, it, "Yes No N/A", "Service twice a week", "", "")
    return _save(d)


def test_THE_bicl_layout_column_header_then_merged_section_rows():
    assert template_docx.read(_bicl_like()) == [
        {"title": "Site Boundaries",
         "items": ["Site secured", "Hoardings established where appropriate"]},
        {"title": "Emergency Procedures", "items": ["Evacuation plan displayed"]},
    ]


def test_the_penina_layout_section_title_in_the_header_row():
    assert template_docx.read(_penina_like()) == [
        {"title": "General", "items": ["Site Hazard Identification Board display",
                                       "Is the H&S Policy displayed?"]},
        {"title": "Site Amenities",
         "items": ["Is the toilet facility in a clean and sanitary state?"]},
    ]


def test_filled_in_answers_and_comments_are_never_read_as_items():
    got = template_docx.read(_penina_like())
    flat = [i for s in got for i in s["items"]]
    assert "Service twice a week" not in flat


def test_a_word_file_with_no_checklist_table_reads_as_empty():
    d = docx.Document()
    d.add_paragraph("Minutes of the site meeting")
    t = d.add_table(rows=0, cols=2)
    _row(t, "Attendees", "Ben, Neil")
    assert template_docx.read(_save(d)) == []


def test_items_before_any_heading_go_in_an_untitled_section():
    d = docx.Document()
    t = d.add_table(rows=0, cols=2)
    _row(t, "Gate locked", "Yes No N/A")
    assert template_docx.read(_save(d)) == [{"title": "", "items": ["Gate locked"]}]


def test_a_w14_checkbox_content_control_alone_marks_an_answer_cell():
    from docx.oxml import parse_xml
    d = docx.Document()
    t = d.add_table(rows=0, cols=2)
    _row(t, "Gate locked", "")
    # A bare checkbox content control with no "Yes No N/A" text beside it.
    t.rows[0].cells[1]._tc.append(parse_xml(
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
        '<w:sdtPr><w14:checkbox/></w:sdtPr>'
        '<w:sdtContent><w:p><w:r><w:t>☐</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
    assert template_docx.read(_save(d)) == [{"title": "", "items": ["Gate locked"]}]


@pytest.mark.parametrize("name,n_sections,n_items",[("bicl", 11, 47), ("penina", 19, 115)])
def test_the_real_forms_read_exactly_as_the_ground_truth(name, n_sections, n_items):
    path = REAL[name]
    if not os.path.exists(path):
        pytest.skip("real form not on this machine: " + path)
    truth = json.load(open(FIXTURES / f"truth_{name}.json", encoding="utf-8"))
    got = template_docx.read(path)
    assert got == truth["sections"]
    assert len(got) == n_sections
    assert sum(len(s["items"]) for s in got) == n_items
