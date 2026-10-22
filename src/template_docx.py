"""Read a checklist out of a Word form -- python-docx, no model (spec 2026-10-16 section 3).

Free, exact and deterministic. The rule (the spike's truth reader, generalised -- ruling R8):
  * an ITEM row is a row any of whose cells after the first is an answer cell: its text
    is "Yes No N/A" once checkbox glyphs are removed, or it holds a w14 checkbox content
    control. The item text is the first cell's text, whitespace-joined.
  * a SECTION is a non-empty single-cell (merged) row, or the first cell of a header row
    (cell 2 reads "Yes/No/NA" or "COMPLIES ...") -- unless the very next row is itself a
    single-cell section row, in which case the header row is only column headings.
  * tables with no item row at all (project details, signatures, comments) are skipped.
  * ANSWER COLUMNS (prod hand test, 2026-10-22): a table whose column-header row has a
    cell of its own for Yes and for No (N/A optional) -- "# | Item | Yes | No | N/A" --
    is read by column instead: an item row has text left of the answer columns and only
    blank or ticked boxes in them; the item text is the longest of those left cells that
    is not just a number. Its section is the header row's own label when that label names
    something ("Scaffolding | Y | N | NA"), else the paragraph just above the table (the
    usual heading), else untitled; merged single-cell rows inside it are sections as above.
The python-docx layer is cp312 (BUG-24); without it AVAILABLE is False and the worker
answers "upload a PDF" instead of raising ImportError.
"""
import logging
import re

logger = logging.getLogger(__name__)

try:
    import docx  # python-docx, from the docx layer
    AVAILABLE = True
except ImportError:  # pragma: no cover - exercised by the worker test via monkeypatch
    docx = None
    AVAILABLE = False

_ANSWER = re.compile(r"^yes\s*no\s*n/?a$", re.IGNORECASE)
_HEADER = re.compile(r"(yes\s*/\s*no|complies)", re.IGNORECASE)
_BOXES = re.compile("[\u2610\u2611\u2612\u2713\u2714\u25a1\u25a0]")
_W14_CHECKBOX = "{http://schemas.microsoft.com/office/word/2010/wordml}checkbox"
_YES, _NO, _NA = {"yes", "y"}, {"no", "n"}, {"n/a", "na", "n.a.", "not applicable"}
# What a ticked box holds once _BOXES glyphs are gone; anything longer is a written answer.
_TICK = re.compile(r"^(x|y|n|yes|no|n/?a)?$", re.IGNORECASE)
_NUMBER = re.compile(r"^[\d.()]*$")
# Column labels that name no section ("Item | Yes | No"): the heading above names it.
_GENERIC = {"", "#", "no", "no.", "item", "items", "description", "activity", "check",
            "checks", "question", "questions", "requirement", "requirements", "task"}
_MAX_TITLE = 100


def _unique_cells(row):
    """A merged cell appears once per grid column it spans; keep it once."""
    out, prev = [], None
    for cell in row.cells:
        if cell._tc is prev:
            continue
        prev = cell._tc
        out.append(cell)
    return out


def _text(cell):
    return " ".join(_BOXES.sub(" ", cell.text).split())


def _is_answer_cell(cell):
    if _ANSWER.match(_text(cell)):
        return True
    return any(el.tag == _W14_CHECKBOX for el in cell._tc.iter())


def _kind(cells):
    """('item'|'header'|'section'|None, first-cell text)."""
    first = _text(cells[0]) if cells else ""
    if len(cells) > 1 and any(_is_answer_cell(c) for c in cells[1:]):
        return "item", first
    if len(cells) > 1 and _HEADER.search(_text(cells[1])):
        return "header", first
    if len(cells) == 1 and first:
        return "section", first
    return None, first


def _label(text):
    return " ".join(text.lower().split())


def _answer_columns(cells):
    """Indexes of the Yes / No / N/A cells when `cells` is a column-header row with a
    cell of its own for at least Yes and No; else None."""
    found = {}
    for i, c in enumerate(cells):
        label = _label(_text(c))
        for kind, words in (("yes", _YES), ("no", _NO), ("na", _NA)):
            if label in words and kind not in found:
                found[kind] = i
    if "yes" not in found or "no" not in found:
        return None
    return sorted(found.values())


def _item_text(cells):
    texts = [_text(c) for c in cells]
    texts = [t for t in texts if t and not _NUMBER.match(t)]
    return max(texts, key=len) if texts else ""


def _read_by_columns(rows, heading):
    """(title, [(kind, text)]) for an answer-column table, or None when it is not one."""
    for h, cells in enumerate(rows):
        cols = _answer_columns(cells)
        if cols is not None:
            break
    else:
        return None
    left = _item_text(rows[h][:cols[0]])
    title = left if _label(left) not in _GENERIC else heading
    out = []
    for cells in rows[h + 1:]:
        if len(cells) == 1:
            if _text(cells[0]):
                out.append(("section", _text(cells[0])))
            continue
        if len(cells) <= cols[-1]:
            continue
        text = _item_text(cells[:cols[0]])
        answers = [cells[i] for i in cols]
        if text and all(_TICK.match(_text(c)) or _is_answer_cell(c) for c in answers):
            out.append(("item", text))
    return title, out


def _headings(document):
    """The text of the paragraph just above each top-level table ("" when another table,
    or nothing, is directly above it), in document.tables order."""
    out, last = [], ""
    for el in document.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = " ".join("".join(t.text or "" for t in el.iter() if t.tag.endswith("}t")).split())
            if text:
                last = text
        elif tag == "tbl":
            out.append(last if len(last) <= _MAX_TITLE else "")
            last = ""
    return out


def read(source):
    """[{title, items: [str]}] in document order; [] when no checklist table is found."""
    document = docx.Document(source)
    sections = []
    headings = _headings(document)
    for t, table in enumerate(document.tables):
        cells = [_unique_cells(r) for r in table.rows]
        rows = [_kind(c) for c in cells]
        if not any(kind == "item" for kind, _ in rows):
            by_columns = _read_by_columns(cells, headings[t] if t < len(headings) else "")
            if not by_columns or not any(kind == "item" for kind, _ in by_columns[1]):
                continue
            title, rows = by_columns
            if title:
                sections.append({"title": title, "items": []})
            for kind, text in rows:
                if kind == "section":
                    sections.append({"title": text, "items": []})
                else:
                    if not sections:
                        sections.append({"title": "", "items": []})
                    sections[-1]["items"].append(text)
            continue
        for i, (kind, text) in enumerate(rows):
            if kind == "item":
                if not sections:
                    sections.append({"title": "", "items": []})
                sections[-1]["items"].append(text)
            elif kind == "header":
                next_kind = rows[i + 1][0] if i + 1 < len(rows) else None
                if next_kind != "section" and text:
                    sections.append({"title": text, "items": []})
            elif kind == "section":
                sections.append({"title": text, "items": []})
    return [s for s in sections if s["items"]]
