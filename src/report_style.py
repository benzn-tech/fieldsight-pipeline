"""One look for every Word report FieldSight writes (owner, 2026-10-08).

Every report -- a template report, a spoken check's checklist, a session
report, the nightly daily report -- went out in python-docx's defaults: black
'Table Grid' boxes, Word's stock blue headings, no running header or page
numbers. A customer forwarding one to a client could not tell it from a
document typed in a hurry. This module is the house style, applied by each
renderer so they cannot drift apart:

  * the page: A4, 2 cm margins, a running header (project left, report title
    right) and footer ("Prepared with FieldSight", generated time, "Page X of
    Y") on every page;
  * type: Calibri, dark grey body, navy headings, a thin accent rule under each
    section heading;
  * tables: navy header row in white, light rules, alternate row shading, the
    header repeated on every page, rows kept whole across a page break;
  * key-value tables (Report Details): a shaded label column instead of a
    header row;
  * checklist answers: Yes / No / N/A as coloured, bold cells, so the items
    needing action are found at a glance.

Never raises into a renderer: styling is presentation, and a document that
could not be styled still goes out plain rather than not at all.
"""
import datetime as dt
import logging

logger = logging.getLogger(__name__)

try:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    AVAILABLE = True
except ImportError:          # the layer is missing: renderers already say so
    AVAILABLE = False

FONT = "Calibri"
NAVY = "1F3A5F"
NAVY_LIGHT = "2E5A88"
ACCENT = "F2B705"            # FieldSight yellow, deepened to print well
TEXT = "2B2B2B"
MUTED = "667085"
RULE = "D0D5DD"
ZEBRA = "F5F7FA"
LABEL = "EEF2F6"
ANSWER_STYLE = {             # fill, text
    "yes": ("E3F4E8", "1E7B34"),
    "no": ("FDE7E7", "B42318"),
    "n/a": ("EEF0F3", "5F6B7A"),
}


def _rgb(hexstr):
    return RGBColor.from_string(hexstr)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _border(element_pr, side, size=6, color=RULE, tag="w:pBdr"):
    bdr = element_pr.find(qn(tag))
    if bdr is None:
        bdr = OxmlElement(tag)
        element_pr.append(bdr)
    edge = OxmlElement("w:%s" % side)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "4")
    edge.set(qn("w:color"), color)
    bdr.append(edge)


def _field(run, code):
    """A Word field (PAGE, NUMPAGES) in `run`."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _font(style, size, color=TEXT, bold=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = _rgb(color)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)


def setup(doc, title, header_left=None, generated=None):
    """Page, type and the running header and footer. Call once, first."""
    if not AVAILABLE:
        return
    try:
        for sec in doc.sections:
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
            sec.left_margin = sec.right_margin = Cm(2.0)
            sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
            sec.header_distance = sec.footer_distance = Cm(1.0)
        styles = doc.styles
        normal = styles["Normal"]
        _font(normal, 10)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.15
        _font(styles["Title"], 22, NAVY, bold=True)
        styles["Title"].paragraph_format.space_after = Pt(2)
        for name, size, color, before in (("Heading 1", 13, NAVY, 16), ("Heading 2", 11, NAVY_LIGHT, 10),
                                          ("Heading 3", 10, MUTED, 8)):
            st = styles[name]
            _font(st, size, color, bold=True)
            st.font.italic = False
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(4)
            st.paragraph_format.keep_with_next = True
        # A thin accent rule under every section heading.
        _border(styles["Heading 1"].element.get_or_add_pPr(), "bottom", 8, ACCENT)
        for name in ("List Bullet",):
            if name in [s.name for s in styles]:
                _font(styles[name], 10)
        _header_footer(doc, title, header_left, generated)
        doc.core_properties.title = title or ""
        doc.core_properties.author = "FieldSight"
    except Exception:
        logger.warning("report style: page setup failed -- the document goes out plain",
                       exc_info=True)


def _header_footer(doc, title, header_left, generated):
    sec = doc.sections[0]
    usable = sec.page_width - sec.left_margin - sec.right_margin
    # Word's own Header/Footer styles carry centre and right tab stops for a
    # Letter page; left in place, "Page 1 of 5" stopped at the centre one.
    for name in ("Header", "Footer"):
        try:
            doc.styles[name].paragraph_format.tab_stops.clear_all()
        except KeyError:
            pass

    head = sec.header.paragraphs[0]
    head.text = ""
    head.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    left = head.add_run(header_left or "FieldSight")
    left.bold = True
    head.add_run("\t")
    right = head.add_run(title or "")
    for r in (left, right):
        r.font.size, r.font.color.rgb, r.font.name = Pt(8), _rgb(MUTED), FONT
    left.font.color.rgb = _rgb(NAVY)
    _border(head._p.get_or_add_pPr(), "bottom", 4, RULE)

    foot = sec.footer.paragraphs[0]
    foot.text = ""
    foot.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    if generated is None:
        try:
            import nz_time
            generated = nz_time.nz_now()          # Lambda's clock is UTC
        except Exception:
            generated = dt.datetime.now()
    when = generated
    a = foot.add_run("Prepared with FieldSight  ·  generated %s" % when.strftime("%d %b %Y %H:%M"))
    foot.add_run("\t")
    b = foot.add_run("Page ")
    c = foot.add_run()
    _field(c, "PAGE")
    d = foot.add_run(" of ")
    e = foot.add_run()
    _field(e, "NUMPAGES")
    for r in (a, b, c, d, e):
        r.font.size, r.font.color.rgb, r.font.name = Pt(8), _rgb(MUTED), FONT
    _border(foot._p.get_or_add_pPr(), "top", 4, RULE)


def title_block(doc, title, subtitle=None):
    """The report's own title on page one, its subtitle, and an accent bar."""
    if not AVAILABLE:
        return doc.add_heading(title, level=0)
    t = doc.add_paragraph(title, style="Title")
    if subtitle:
        p = doc.add_paragraph()
        run = p.add_run(subtitle)
        run.font.size, run.font.color.rgb = Pt(11), _rgb(MUTED)
        p.paragraph_format.space_after = Pt(2)
        bar = p
    else:
        bar = t
    try:
        _border(bar._p.get_or_add_pPr(), "bottom", 18, ACCENT)
        bar.paragraph_format.space_after = Pt(12)
    except Exception:
        logger.debug("title bar not drawn", exc_info=True)
    return t


def _tbl_borders(table):
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:%s" % side)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), RULE)
        borders.append(edge)
    tblpr.append(borders)


def _cell_runs(cell):
    return [r for p in cell.paragraphs for r in p.runs]


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _keep_whole(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trpr.append(el)


NOT_COVERED = "Not covered in the recording"     # checklist.NOT_COVERED
# A checklist's columns, as shares of the text width: the item is the
# sentence a reader scans; Yes / No needs a word.
CHECKLIST_SHARES = {"item": 0.36, "answer": 0.09, "comment": 0.31, "responsible": 0.12,
                    "due": 0.12, "photos": 0.16}


def _widths(table, heads):
    """Column widths for a checklist table, from CHECKLIST_SHARES."""
    shares = [CHECKLIST_SHARES.get(h, 0.12) for h in heads]
    total = sum(shares)
    usable = Cm(17.0)
    table.autofit = False
    for row in table.rows:
        for cell, share in zip(row.cells, shares):
            cell.width = int(usable * share / total)


_MARK = "fieldsight-styled"


def _styled(table):
    cap = table._tbl.tblPr.find(qn("w:tblCaption"))
    return cap is not None and cap.get(qn("w:val")) == _MARK


def _mark(table):
    cap = OxmlElement("w:tblCaption")
    cap.set(qn("w:val"), _MARK)
    table._tbl.tblPr.append(cap)


def style_all(doc):
    """Every table not styled yet, with the default (header-row) style. Call
    once, last, so a renderer only has to name the special ones."""
    if not AVAILABLE:
        return
    for table in doc.tables:
        if not _styled(table):
            style_table(table)


def style_table(table, facts=False):
    """The house table. `facts`: a key-value table (Report Details) -- label
    column shaded, no header row. Otherwise row 0 is the header. A column
    headed "Answer" gets its Yes / No / N/A cells coloured."""
    if not AVAILABLE:
        return
    try:
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _tbl_borders(table)
        rows = table.rows
        answer_col = None
        if not facts and rows:
            heads = [c.text.strip().lower() for c in rows[0].cells]
            answer_col = heads.index("answer") if "answer" in heads else None
            if answer_col is not None:
                _widths(table, heads)
        for r, row in enumerate(rows):
            _keep_whole(row)
            for c, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                for run in _cell_runs(cell):
                    run.font.size, run.font.name = Pt(9), FONT
                if facts:
                    if c == 0:
                        _shade(cell, LABEL)
                        for run in _cell_runs(cell):
                            run.bold, run.font.color.rgb = True, _rgb(NAVY)
                    continue
                if r == 0:
                    _shade(cell, NAVY)
                    for run in _cell_runs(cell):
                        run.bold, run.font.color.rgb = True, _rgb("FFFFFF")
                    continue
                if r % 2 == 0:
                    _shade(cell, ZEBRA)
                if cell.text.strip() == NOT_COVERED:
                    for run in _cell_runs(cell):
                        run.italic, run.font.color.rgb = True, _rgb(MUTED)
                if c == answer_col:
                    style = ANSWER_STYLE.get(cell.text.strip().lower())
                    if style:
                        _shade(cell, style[0])
                        for run in _cell_runs(cell):
                            run.bold, run.font.color.rgb = True, _rgb(style[1])
                        for p in cell.paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not facts and rows:
            _repeat_header(rows[0])
        _mark(table)
    except Exception:
        logger.warning("report style: a table could not be styled -- left plain", exc_info=True)


def caption(doc, text):
    """A small grey caption under a photograph strip."""
    if not AVAILABLE or not text:
        return
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size, run.font.color.rgb = Pt(8), _rgb(MUTED)
    p.paragraph_format.space_after = Pt(6)
